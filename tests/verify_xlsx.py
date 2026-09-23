"""Validate generated OOXML entirely in memory; no spreadsheet fixture files."""
import io
import sys
import zipfile
import xml.etree.ElementTree as ET
from openpyxl import load_workbook
payload = sys.stdin.buffer.read()
with zipfile.ZipFile(io.BytesIO(payload)) as z:
    assert z.testzip() is None
    for name in z.namelist():
        ET.fromstring(z.read(name))
        if name.startswith('xl/worksheets/'):
            assert b'<f>' not in z.read(name)
wb = load_workbook(io.BytesIO(payload))
assert wb.sheetnames[0] == '전체'
assert len(wb.sheetnames) >= 2
assert len({s.lower() for s in wb.sheetnames}) == len(wb.sheetnames)
for ws in wb:
    assert ws.max_column == 8
    assert [ws.cell(4,c).value for c in range(1,9)] == ['점검일','동','코어','호수','위치','공종','내용','확인사항']
    assert ws.freeze_panes == 'A5'
    assert ws.page_setup.fitToWidth == 1
    assert ws.page_setup.orientation == 'landscape'
    assert ws.auto_filter.ref == f'A4:H{ws.max_row}'
    assert 'A1:H1' in ws.merged_cells
    for row in ws.iter_rows(min_row=5):
        for cell in row:
            assert cell.data_type != 'f'
if '--fixture' in sys.argv:
    ws=wb['전체']
    assert [ws.cell(i,4).value for i in range(5,9)] == ['1501','1401','1503','1401']
    assert ws['G5'].value == ' =1+1\t줄\n보수 <&> "따옴표"'
    assert ws['A5'].value == '날짜 미상'
    assert ws['B2'].value == '샘플현장'
    assert ws['G2'].value == '점검자 A'
if '--review' in sys.argv:
    ws=wb['전체']
    assert [ws.cell(i,8).value or '' for i in range(5,8)] == ['공종 확인필요','','']
    assert ws['H5'].fill.fgColor.rgb == 'FFFFE8A3'
    assert wb['공종 확인필요'].max_row == 5
    assert wb['공종 확인필요']['G5'].value == '우측 벽 균열'
    assert '공종 확인필요 (2)' in wb.sheetnames
    assert sum(sheet.max_row-4 for sheet in list(wb)[1:]) == 3
    import importlib.util
    from pathlib import Path
    spec=importlib.util.spec_from_file_location('reference_export',Path(__file__).resolve().parents[1]/'scripts/build_punch_xlsx.py')
    ref=importlib.util.module_from_spec(spec);spec.loader.exec_module(ref)
    from openpyxl import Workbook
    reference=Workbook();reference.remove(reference.active)
    sheet=ref.make_sheet(reference,'전체',[{'공종':'','내용':'우측 벽 균열'}],'샘플현장','','날짜 미상')
    assert sheet['H5'].value == '공종 확인필요' and sheet.max_column == 8
    assert sheet.auto_filter.ref == 'A4:H5'
print('ZIP CRC, XML, openpyxl, sheets, headers, string cells, print layout: OK')
