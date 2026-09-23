# -*- coding: utf-8 -*-
"""
펀치리스트 엑셀 만들기 — 현장에서 쓰던 CHECK LIST 양식 그대로.

일반적인 세대 점검 양식의 열 구성을 따르되 입력을 줄였다.
  · 뺀 것: 세부공종 · 유형 (말이 길어져 음성 입력이 번거로워진다)
  · 뺀 것: 완료일 · 확인 · 개수 · 전달일 · 보수완료율 (실제로 비어 있었고, 받은 뒤 각자 붙여 쓴다)
공종은 현장에서 부르는 단위 그대로 넣는다 — 할석 · 견출 · 미장 · 석고 · 도배 …

  전체 시트 하나 + 공종별 시트 (협력사에 나눠 줄 때 그대로 인쇄)
  정렬: 동 → 코어 → 층 내림차순  (코어 하나를 위층부터 내려오며 보수하도록)

사용법:
    python scripts/build_punch_xlsx.py --out data/펀치리스트_샘플.xlsx --sample
"""
import argparse
import json
import os
import re

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DATA = os.path.join(ROOT, "data")

HEADERS = ["점검일", "동", "코어", "호수", "위치", "공종", "내용", "확인사항"]
WIDTHS = [11, 7, 9, 8, 13, 11, 62, 20]

NAVY = "0F3557"
THIN = Side(style="thin", color="B7BFC8")
MED = Side(style="medium", color=NAVY)
BODY_FONT = Font(name="맑은 고딕", size=10)
HEAD_FONT = Font(name="맑은 고딕", size=10, bold=True, color="FFFFFF")
LABEL_FONT = Font(name="맑은 고딕", size=10, bold=True)
TITLE_FONT = Font(name="맑은 고딕", size=16, bold=True, color=NAVY)


def num_of(v):
    m = re.search(r"(\d+)", str(v))
    return int(m.group(1)) if m else 0


def floor_of(ho):
    """호수에서 층을 뽑는다. 1501 → 15, 901 → 9"""
    n = num_of(ho)
    return n // 100 if n >= 100 else n


def parse_cores(text):
    """현장 세팅 문자열을 코어 구성으로 바꾼다.

    "1코어=01,02; 2코어=03,04"  →  {"01": "1코어", "02": "1코어", ...}
    라인은 호수 뒤 두 자리다 (1503 → 03).  현장마다 다르므로 사용자가 정한다.
    """
    table = {}
    for part in str(text).split(";"):
        if "=" not in part:
            continue
        core, lines = part.split("=", 1)
        for ln in lines.split(","):
            ln = ln.strip()
            if ln:
                table["%02d" % int(re.sub(r"\D", "", ln))] = core.strip()
    return table


def line_of(ho):
    """호수에서 라인을 뽑는다. 1503 → '03', 901 → '01'"""
    n = num_of(ho)
    return "%02d" % (n % 100) if n >= 100 else "%02d" % n


def fill_cores(rows, table):
    """코어 칸이 비어 있으면 호수를 보고 채운다."""
    if not table:
        return rows
    for r in rows:
        if not r.get("코어") and r.get("호수"):
            r["코어"] = table.get(line_of(r["호수"]), "")
    return rows


def sort_rows(rows):
    """동 → 코어 → 층 내림차순 → 호수.

    작업자가 코어 하나를 잡고 위층부터 내려오며 보수하는 순서 그대로다.
    """
    return sorted(rows, key=lambda r: (str(r.get("동", "")), num_of(r.get("코어", "")),
                                       -floor_of(r.get("호수", "")), num_of(r.get("호수", ""))))


def make_sheet(wb, title, rows, site, inspector, date_text):
    ws = wb.create_sheet(title[:31])
    last = get_column_letter(len(HEADERS))

    ws.merge_cells("A1:%s1" % last)
    ws["A1"] = "CHECK LIST"
    ws["A1"].font = TITLE_FONT
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 28

    # 머리 — 현장명·점검일·점검인은 채우고, 배포할 때 손으로 쓰는 칸은 비워 둔다
    labels = {"A2": "현장명", "D2": "점검일", "F2": "점검인",
              "A3": "완료요구일", "D3": "수령인", "F3": "연락처"}
    values = {"B2": site, "E2": date_text, "G2": inspector,
              "B3": "", "E3": "", "G3": ""}
    for ref, val in labels.items():
        ws[ref] = val
        ws[ref].font = LABEL_FONT
        ws[ref].fill = PatternFill("solid", fgColor="EDF1F5")
        ws[ref].alignment = Alignment(horizontal="center", vertical="center")
    for ref, val in values.items():
        ws[ref] = val
        ws[ref].font = BODY_FONT
        ws[ref].alignment = Alignment(horizontal="left", vertical="center")
    ws.merge_cells("B2:C2")
    ws.merge_cells("B3:C3")
    ws.merge_cells("G2:H2")
    ws.merge_cells("G3:H3")
    for r in (2, 3):
        ws.row_dimensions[r].height = 20

    for i, h in enumerate(HEADERS, 1):
        c = ws.cell(row=4, column=i, value=h)
        c.font = HEAD_FONT
        c.fill = PatternFill("solid", fgColor=NAVY)
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = Border(top=MED, bottom=MED, left=THIN, right=THIN)
        ws.column_dimensions[get_column_letter(i)].width = WIDTHS[i - 1]
    ws.row_dimensions[4].height = 22

    for n, row in enumerate(sort_rows(rows)):
        r = 5 + n
        for i, key in enumerate(HEADERS, 1):
            value = ("공종 확인필요" if not str(row.get("공종") or "").strip() else "") if key == "확인사항" else row.get(key, "")
            c = ws.cell(row=r, column=i, value=value)
            c.data_type = "s"
            if key == "확인사항" and value:
                c.fill = PatternFill("solid", fgColor="FFFFE8A3")
            c.font = BODY_FONT
            c.border = Border(top=THIN, bottom=THIN, left=THIN, right=THIN)
            c.alignment = Alignment(
                horizontal="left" if key == "내용" else "center",
                vertical="center", wrap_text=(key in ("내용", "확인사항")))
        # 내용이 길면 줄바꿈되도록 행 높이는 지정하지 않는다 (엑셀 자동 맞춤)

    end = 4 + len(rows)
    ws.auto_filter.ref = "A4:%s%d" % (last, end)
    ws.freeze_panes = "A5"

    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = "4:4"
    ws.print_options.horizontalCentered = True
    ws.page_margins.left = ws.page_margins.right = 0.3
    ws.page_margins.top = ws.page_margins.bottom = 0.5
    return ws


def build(rows, out, site, inspector, date_text):
    wb = Workbook()
    wb.remove(wb.active)
    make_sheet(wb, "전체", rows, site, inspector, date_text)

    pending = [r for r in rows if not str(r.get("공종") or "").strip()]
    trades = []
    used = {"전체", "공종 확인필요", "history"}
    if pending:
        make_sheet(wb, "공종 확인필요", pending, site, inspector, date_text)
        trades.append("공종 확인필요")
    keys = list(dict.fromkeys(r["공종"] for r in rows if str(r.get("공종") or "").strip()))
    for trade in keys:
        base = re.sub(r"[\[\]:*?/\\\x00-\x1f]", "·", trade).strip("'").strip()[:31] or "미지정"
        name, suffix = base, 2
        while name.lower() in used:
            tail = " (%d)" % suffix
            name = base[:31-len(tail)] + tail
            suffix += 1
        used.add(name.lower())
        make_sheet(wb, name, [r for r in rows if r.get("공종") == trade], site, inspector, date_text)
        trades.append(name)

    os.makedirs(os.path.dirname(out) or ".", exist_ok=True)
    wb.save(out)
    return out, trades


def trade_word(p):
    """공종은 현장에서 부르는 대로 적는다.

    골조는 '골조할석 · 골조미장'처럼 붙여 부르고, 내장은 '합지 · 경량석고'처럼
    세부만 부르기도 한다. 어느 쪽이든 부른 그대로 한 칸에 들어간다.
    """
    big, sub = p["공종"], (p["세부공종"] or "").replace("/", "")
    if big == "골조" and sub:
        return "골조" + {"할미": "할석", "견출": "견출", "미장": "미장"}.get(sub, sub)
    return sub or big


# ---------------------------------------------------------------- 견본 자료
def sample_rows():
    """뽑아 둔 문구 사전으로 견본 펀치리스트를 만든다."""
    path = os.path.join(DATA, "punch_terms.json")
    with open(path, encoding="utf-8") as fh:
        terms = json.load(fh)
    phrases = [p for p in terms["문구"] if p["횟수"] >= 2 and p["공종"]][:22]

    date_text = "26.09.16"
    rows = []
    i = 0
    # 1코어는 01·02호, 2코어는 03·04호.  코어 하나를 위층부터 훑고 다음 코어로 넘어간다.
    for dong in ("101", "102"):
        for units in ((1, 2), (3, 4)):
            for floor in (15, 14, 13):
                for unit in units:
                    for _ in range(1 if (floor + unit) % 2 else 2):
                        p = phrases[i % len(phrases)]
                        i += 1
                        spot = (p["주로 쓰는 위치"].split(" · ")[0]
                                if p["주로 쓰는 위치"] else "거실")
                        rows.append({
                            "점검일": date_text, "동": dong, "코어": "",
                            "호수": floor * 100 + unit, "위치": spot,
                            "공종": trade_word(p), "내용": p["문구"],
                        })
    return rows, date_text


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--out", default=os.path.join(DATA, "펀치리스트_샘플.xlsx"))
    ap.add_argument("--site", default="샘플현장")
    ap.add_argument("--inspector", default="")
    ap.add_argument("--cores", default="1코어=01,02; 2코어=03,04",
                    help="현장 세팅 — 코어별 라인 (예: '1코어=01,02; 2코어=03,04')")
    ap.add_argument("--sample", action="store_true", help="견본 자료로 만들기")
    a = ap.parse_args()

    if not a.sample:
        ap.error("지금은 --sample 로 견본만 만듭니다. (시스템 연결은 나중에)")

    rows, date_text = sample_rows()
    rows = fill_cores(rows, parse_cores(a.cores))
    out, trades = build(rows, a.out, a.site, a.inspector, date_text)
    print("%d건 · 시트 %d개 (전체 + %s)" % (len(rows), len(trades) + 1, " · ".join(trades)))
    print(out)


if __name__ == "__main__":
    main()
